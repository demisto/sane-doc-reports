from docx import Document
from docx.shared import Pt, Mm

from sane_doc_reports.CellObject import CellObject
from sane_doc_reports.Section import sane_to_section, Section
from sane_doc_reports.utils import insert_by_type
from sane_doc_reports.conf import DEBUG, LAYOUT_KEY, STYLE_KEY, \
    A4_MM_HEIGHT, A4_MM_WIDTH, TOP_MARGIN_PT, BOTTOM_MARGIN_PT, \
    LEFT_MARGIN_PT, RIGHT_MARGIN_PT
from sane_doc_reports.SaneJson import SaneJson
from sane_doc_reports.grid import get_cell, merge_cells


class Report:
    """
    In charge of generating a DOCX report form a SANE report (JSON)
    """

    def __init__(self, json_file_path: str):
        self.document = Document()
        self.sane_json = SaneJson(json_file_path)

    def populate_report(self) -> None:
        self.change_page_size('A4')
        self._decrease_layout_margins()
        for page_num, page in enumerate(self.sane_json.get_pages()):
            cols, rows = page.calculate_page_grid()

            if DEBUG:
                print(f'Creating a layout grid of size ({rows},{cols})' +
                      f' for page: {page_num}')
            grid = self.document.add_table(rows=rows, cols=cols)

            if DEBUG:
                grid.style = 'Table Grid'

            for section in page.get_sections():
                cell = get_cell(grid, section)
                merge_cells(grid, section)
                cell_object = CellObject(cell)
                section = sane_to_section(section)

                self._insert_section(cell_object, section)

    @staticmethod
    def _insert_section(cell_object: CellObject, section: Section) -> None:
        section_type = section.type

        # Fix the chart name
        if section_type == 'chart':
            section_type = section.layout['chartType'] + '_chart'
            section.type = section_type

        insert_by_type(section_type, cell_object, section)

    def save(self, output_file_path: str):
        self.document.save(output_file_path)

    def change_page_size(self, paper_size: str, direction=None) -> None:
        if paper_size == 'A4':
            sections = self.document.sections
            for section in sections:
                section.page_height = Mm(A4_MM_HEIGHT)
                section.page_width = Mm(A4_MM_WIDTH)

    def _decrease_layout_margins(self) -> None:
        sections = self.document.sections
        for section in sections:
            section.top_margin = Pt(TOP_MARGIN_PT)
            section.bottom_margin = Pt(BOTTOM_MARGIN_PT)
            section.left_margin = Pt(LEFT_MARGIN_PT)
            section.right_margin = Pt(RIGHT_MARGIN_PT)
