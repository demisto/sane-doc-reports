from typing import Tuple, List

from docx.table import Table, _Cell
from docx.text.paragraph import Paragraph
from docx.text.run import Run

from sane_doc_reports.positioning import row_pos, col_pos, get_height, get_width


def _merge_horizontally(grid: Table, row: int, col: int, width: int) -> None:
    """
    Merge the base cell at row->col with each adjacent cells in the width range.
    """
    for cell_index in range(1, width):
        cell_to_merge = grid.rows[row].cells[col + cell_index]
        grid.rows[row].cells[col].merge(cell_to_merge)


def _merge_vertically(grid: Table, row: int, col: int, height: int,
                      require_second_horizontal_merge: bool,
                      width: int) -> None:
    """
    Merge the base cell provided at row->col.
    Merges with each adjacent cells in the width range (width x cells
    to the right), here we need to merge horizontally before merging vertically
    (because it won't enable us to merge ("span not rectangular"))
    """

    for row_index in range(1, height):
        if require_second_horizontal_merge:
            _merge_horizontally(grid, row + row_index, col, width)

        cell_to_merge = grid.rows[row + row_index].cells[col]
        grid.rows[row].cells[col].merge(cell_to_merge)


def merge_cells(grid: Table, section: dict) -> None:
    """
    Merge the sections cell using it's width and height.
    """
    row, col = row_pos(section), col_pos(section)
    width, height = get_width(section), get_height(section)
    require_second_horizontal_merge = False

    # Merge horizontally
    if width > 1:
        _merge_horizontally(grid, row, col, width)
        require_second_horizontal_merge = True

    # Merge Vertically
    if height > 1:
        _merge_vertically(grid, row, col, height,
                          require_second_horizontal_merge, width)


def get_cell(table: Table, section: dict) -> _Cell:
    """
    Get the sections corresponding cell
    """
    row, col = row_pos(section), col_pos(section)
    return table.rows[row].cells[col]


def get_cell_wrappers(cell: _Cell) -> Tuple[Paragraph, Run]:
    """
    Return the cell's paragraph and create a run object too, return them
    both
    """
    paragraphs = cell.paragraphs
    paragraph = paragraphs[0]
    run = paragraph.add_run()
    return paragraph, run


def get_vtable_merged(table: Table) -> List[List]:
    """
    Return a virtual representation of a table, and make merged cells  0
    and normal cells 1.

    Notes: this will put a zero without regard to row span or col span (!)
    Note 2: If you are trying to understand this, good luck! :O
    JK, You should probably create a table via python-docx and try to use this
    on it, also to trigger the `1 not in vtables` you need to fully merge 2
    rows.
    Note 3: ._tc is the original object, it helps us find out if 2 cells are
    the same cells or not (after merging them).
    """
    vtable = []
    last_cells = []

    for row in table.rows:
        # Fix zero rows
        if len(vtable) > 0:
            if 1 not in vtable[-1]:
                del vtable[-1]

        vtable.append([])
        for cell in row.cells:
            # Skip merged
            if cell._tc in last_cells:
                vtable[-1].append(0)
                continue

            vtable[-1].append(1)
            last_cells.append(cell._tc)
    return vtable
